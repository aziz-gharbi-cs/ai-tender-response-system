from pathlib import Path
from uuid import uuid4

from docx import Document

from app.domain.interfaces.document_parser import DocumentParser
from app.domain.models.document import (
    DocumentPage,
    DocumentType,
    NormalizedDocument,
)


class DOCXParser(DocumentParser):

    async def parse(
        self,
        file_path: str,
    ) -> NormalizedDocument:
        path = Path(file_path)

        if not path.exists():
            raise FileNotFoundError(
                f"DOCX file not found: {file_path}"
            )

        if path.suffix.lower() != ".docx":
            raise ValueError(
                f"Expected a DOCX file, got: {path.suffix}"
            )

        document = Document(str(path))

        paragraphs = [
            paragraph.text.strip()
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        content = "\n".join(paragraphs)

        return NormalizedDocument(
            document_id=str(uuid4()),
            filename=path.name,
            document_type=DocumentType.DOCX,
            pages=[
                DocumentPage(
                    page_number=1,
                    content=content,
                )
            ],
            metadata={
                "source_path": str(path),
                "paragraph_count": len(paragraphs),
            },
        )
